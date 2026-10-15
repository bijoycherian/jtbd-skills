#!/usr/bin/env python3
"""ODI opportunity scoring + JTBD opportunity landscape.

Usage:
    python landscape.py STATEMENTS.csv|.xlsx [--ratings RATINGS.csv] [--combined]
                        [--out DIR] [--name NAME] [--title PREFIX] [--assumed]

Two ways to supply ratings:
  1. Aggregated: importance/satisfaction (% rating 4-5) and n columns in STATEMENTS.
  2. Raw (preferred): --ratings with respondent-level 1-5 answers, either
       long: respondent_id, persona, label, importance, satisfaction (one row per job), or
       wide: respondent_id, persona, attention_pass, J01_imp, J01_sat, ... (survey kit export)
     NA or blank = missing. Rows failing the attention check, giving one answer to every
     question, or (with duration_sec) finishing in under a third of the median are dropped.
     Each job is scored on its own persona's respondents (--pool: everyone).

Input columns (case-insensitive; only importance and satisfaction are required):
    label|id, segment, persona, stage, context, barrier, goal, outcome,
    importance, satisfaction, n, evidence

Writes to DIR:
    NAME_map.html  interactive page: live chart + editable scores (--html)
    NAME_scores.xlsx     scores with live formulas, Evidence, Survey and Rules sheets
    NAME_ranked.csv      scores, zones and review flags
    NAME_map.png         static chart (one view of all personas; --by-persona for one per persona)

Extra columns used when present: source (ai | evidence_update | judgement | survey;
v0.6 'edited' reads as judgement), confidence,
importance_reason, satisfaction_reason, evidence and contradicts
("ID: excerpt || ID: excerpt").

Zone rules follow the team guide; see references/scoring.md.
Optional dependencies: openpyxl (xlsx), matplotlib (charts). Without them,
the CSV is still written.
"""
import argparse, csv, math, os, random, re, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import jtbd_files as jf  # noqa: E402  shared contract: score sources, decision log, Based on line

# ---- Zone rules (team guide). Change cut-offs here only. --------------------
TABLE_STAKES_SAT = 85      # satisfaction above this -> Table stakes
MODEST_SAT_MIN = 80        # satisfaction 80..85 ...
MODEST_IMP_MIN = 90        # ... and importance >= 90 -> Modest opportunity
SOLID_OPP = 10             # opportunity 10..12 with satisfaction < 80 -> Solid
HIGH_OPP = 12              # opportunity > 12 with satisfaction < 80 -> High
MIN_N = 30                 # fewer respondents per persona -> "small sample" flag
# ----------------------------------------------------------------------------

VERSION = "1.0.1"

ZONES = ["High opportunity", "Solid opportunity", "Modest opportunity",
         "Table stakes", "Served right", "Over-served"]


def opportunity(imp, sat):
    return imp / 10 + max(imp - sat, 0) / 10


def zone(imp, sat):
    opp = opportunity(imp, sat)
    if sat > imp:
        return "Over-served"
    if sat > TABLE_STAKES_SAT:
        return "Table stakes"
    if sat >= MODEST_SAT_MIN and imp >= MODEST_IMP_MIN:
        return "Modest opportunity"
    if sat < MODEST_SAT_MIN and opp > HIGH_OPP:
        return "High opportunity"
    if sat < MODEST_SAT_MIN and opp >= SOLID_OPP:
        return "Solid opportunity"
    return "Served right"


STABLE = 80                # % of resamples that must land in the same zone


def stability(imp, sat, n, z, draws=2000, n_sat=None):
    """% of binomial resamples (same n) that land in zone z."""
    n_sat = n_sat or n
    rng = random.Random(f"{imp}-{sat}-{n}-{n_sat}")
    hit = 0
    for _ in range(draws):
        i = 100 * sum(rng.random() < imp / 100 for _ in range(n)) / n
        s_ = 100 * sum(rng.random() < sat / 100 for _ in range(n_sat)) / n_sat
        hit += zone(i, s_) == z
    return round(100 * hit / draws)


def margin(p, n):
    p = min(max(p / 100, 0), 1)
    return 1.96 * math.sqrt(p * (1 - p) / n) * 100


# ---- Review heuristics (prompts for a human, not verdicts) -----------------
NOT_EXIST = re.compile(r"\b(don'?t|do not|can'?t|cannot|no) (have )?(the )?(option|way|ability|access)\b|\bcan'?t\b|\bcannot\b", re.I)
SOLUTION = re.compile(r"\b(app|feature|button|chat|dashboard|portal|integration|notification|reminders?|"
                      r"automatic(ally)?|real[- ]time|documentation|api|plugin|widget|screen|tab)\b", re.I)
# Company voice: internal diagnoses and product jargon. Add your own acronyms here.
COMPANY = re.compile(r"\b(because (we|the company|our product) (is|are)n?'?t|our platform|marketplace|PDP|SKU|e2e)\b", re.I)


def read_statements(path):
    """Statement text only (ratings not required)."""
    if path.lower().endswith((".xlsx", ".xlsm")):
        import openpyxl
        ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        hi = next(i for i, r in enumerate(rows) if any(c and str(c).lower() in ("label", "id") for c in r))
        hdr = [str(c) if c else None for c in rows[hi]]
        it = (dict(zip(hdr, r)) for r in rows[hi + 1:])
    else:
        it = csv.DictReader(open(path, newline="", encoding="utf-8-sig"))
    out = []
    for r in it:
        d = {("label" if k.strip().lower() == "id" else k.strip().lower()): ("" if v is None else str(v).strip())
             for k, v in r.items() if k}
        if d.get("label") and (d.get("record_type") or "job") == "job":
            out.append(d)
    return out


def norm(rows):
    alias = {"id": "label", "importance (a)": "importance", "satisfaction (a)": "satisfaction"}
    out = []
    for r in rows:
        d = {}
        for k, v in r.items():
            if k is None:
                continue
            k2 = alias.get(k.strip().lower(), k.strip().lower())
            d[k2] = "" if v is None else str(v).strip()
        if (d.get("record_type") or "job") != "job":
            continue
        if d.get("importance") and d.get("satisfaction"):
            if "source" in d:
                d["source"] = jf.norm_source(d["source"])
            out.append(d)
    return out


def read_logs(path):
    """Decision log rows (record_type = log) from a schema 1.0 jobs file; [] otherwise."""
    if not path.lower().endswith(".csv"):
        return []
    try:
        return jf.read_jobs(path)[1]
    except (OSError, KeyError):
        return []


def read(path):
    if path.lower().endswith((".xlsx", ".xlsm")):
        import openpyxl
        ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        # skip banner rows until a header containing importance
        hi = next(i for i, r in enumerate(rows)
                  if any(c and "importance" in str(c).lower() for c in r))
        hdr = [str(c) if c else None for c in rows[hi]]
        return norm(dict(zip(hdr, r)) for r in rows[hi + 1:])
    with open(path, newline="", encoding="utf-8-sig") as f:
        return norm(csv.DictReader(f))


NA = {"", "na", "n/a", "doesn't apply", "doesn\u2019t apply", "-"}
FAIL = {"no", "n", "false", "0", "fail", "failed"}


def _rating(v):
    v = (v or "").strip().lower()
    if v in NA:
        return None
    try:
        x = float(v)
    except ValueError:
        return None
    return x if 1 <= x <= 5 else None


def read_ratings(path):
    """Respondent-level ratings in either layout, as (respondent, persona, job, imp, sat, row).

    Long:  respondent_id, persona, label, importance, satisfaction   (one row per job)
    Wide:  respondent_id, persona, attention_pass, J01_imp, J01_sat, ...  (survey kit export)
    NA, blank or anything outside 1-5 counts as missing, never as a rating."""
    with open(path, newline="", encoding="utf-8-sig") as f:
        raw = [{k.strip().lower(): (v or "").strip() for k, v in r.items() if k} for r in csv.DictReader(f)]
    if not raw:
        return [], "empty"
    wide = sorted({k[:-4] for k in raw[0] if k.endswith("_imp") and k[:-4] + "_sat" in raw[0]})
    out = []
    for i, r in enumerate(raw):
        rid = r.get("respondent_id") or r.get("respondent") or f"row{i + 2}"
        per = r.get("persona") or r.get("segment") or ""
        if wide:
            for j in wide:
                out.append((rid, per, j, _rating(r[j + "_imp"]), _rating(r[j + "_sat"]), r))
        else:
            out.append((rid, per, r.get("label") or r.get("id") or r.get("job_id"),
                        _rating(r.get("importance")), _rating(r.get("satisfaction")), r))
    return out, "wide" if wide else "long"


def screen(recs):
    """Drop respondents who failed the attention check, straight-lined or sped through."""
    by_r = {}
    for rec in recs:
        by_r.setdefault(rec[0], []).append(rec)
    drop = {"failed attention check": set(), "same answer to every question": set(), "too fast": set()}
    durs = sorted(float(v[0][5]["duration_sec"]) for v in by_r.values()
                  if re.fullmatch(r"\d+(\.\d+)?", v[0][5].get("duration_sec", "")))
    med = durs[len(durs) // 2] if durs else None
    for rid, rs in by_r.items():
        row = rs[0][5]
        if row.get("attention_pass", "").lower() in FAIL:
            drop["failed attention check"].add(rid)
            continue
        ans = [x for rec in rs for x in (rec[3], rec[4]) if x is not None]
        if len(ans) >= 10 and len(set(ans)) == 1:
            drop["same answer to every question"].add(rid)
            continue
        d = row.get("duration_sec", "")
        if med and re.fullmatch(r"\d+(\.\d+)?", d) and float(d) < med / 3:
            drop["too fast"].add(rid)
    gone = set().union(*drop.values())
    return [rec for rec in recs if rec[0] not in gone], {k: len(v) for k, v in drop.items() if v}


def _pct(a, b):
    return round(100 * a / b, 1) if b else 0.0


def aggregate(statements, ratings_path, combined, by_persona=False, pool=False):
    """Turn respondent-level 1-5 ratings into top-2-box % and n.

    Default (one view): each job is scored by respondents of the job's own persona, when the
    job names one and those respondents exist; otherwise by everyone. pool=True scores every
    job on everyone. Importance and satisfaction keep separate bases (NA stays out)."""
    recs, layout = read_ratings(ratings_path)
    recs, dropped = screen(recs)
    if dropped:
        print("Dropped respondents: " + ", ".join(f"{v} {k}" for k, v in dropped.items()))
    by_label = {}
    for s in statements:
        by_label[s["label"]] = s
        if s.get("job_id"):
            by_label[s["job_id"]] = s
    for s in statements:  # J01 -> label 1 when no job_id column
        m = re.fullmatch(r"\d+", s["label"])
        if m:
            by_label.setdefault(f"J{int(s['label']):02d}", s)
    groups, split, unknown = {}, {}, set()

    def add(g, i, sa):
        if i is not None:
            g[0] += 1; g[1] += i >= 4; g[2] += i <= 2
        if sa is not None:
            g[3] += 1; g[4] += sa >= 4; g[5] += sa <= 2

    for rid, per, job, i, sa, _ in recs:
        st = by_label.get(job) or by_label.get(str(job).lower()) or by_label.get(str(job).upper())
        if st is None:
            if i is not None or sa is not None:
                unknown.add(job)
            continue
        lab = st["label"]
        seg = per or st.get("persona") or "All"
        add(split.setdefault(lab, {}).setdefault(seg, [0] * 6), i, sa)
        if by_persona:
            keys = [(seg, lab)] + ([("All (combined)", lab)] if combined and seg != "All" else [])
            for k in keys:
                add(groups.setdefault(k, [0] * 6), i, sa)
    for job in sorted(unknown):
        print(f"warning: ratings for unknown job {job!r} ignored", file=sys.stderr)
    if not by_persona:
        for lab, per in split.items():
            own = by_label[lab].get("persona", "")
            use = [own] if (not pool and own in per) else list(per)
            g = groups.setdefault(("All", lab), [0] * 6)
            for p in use:
                for k in range(6):
                    g[k] += per[p][k]
            if not pool and own and own not in per and len(per) > 1:
                print(f"note: job {lab} names persona {own!r}, but no respondent chose it; scored on everyone")
            by_label[lab]["_scored_on"] = own if len(use) == 1 and own in per else ""
    out = []
    for (seg, lab), (ni, ti, bi, ns, ts, bs) in groups.items():
        if not ni or not ns:
            continue
        d = dict(by_label[lab])
        d.update(segment=seg, n=str(min(ni, ns)), importance=str(_pct(ti, ni)),
                 satisfaction=str(_pct(ts, ns)), _n=(ni, ns),
                 _b2b=(_pct(bi, ni), _pct(bs, ns)))
        if not seg.startswith("All"):
            d["persona"] = seg
        d["source"] = "survey"
        if not by_persona and not d.get("persona"):
            d["persona"] = ", ".join(sorted(split.get(lab, {}))) or "All"
        if not by_persona and len(split.get(lab, {})) > 1:
            d["_split"] = {p: (min(v[0], v[3]), round(_pct(v[1], v[0])), round(_pct(v[4], v[3])))
                           for p, v in split[lab].items() if v[0] and v[3]}
        out.append(d)
    return out


def score(rows, assumed_flag, by_persona=False):
    has_n = all(r.get("n") or (r.get("source") or "").lower() == "ai" for r in rows)
    ai = any((r.get("source") or "").lower() == "ai" for r in rows)
    assumed = assumed_flag or ai or not has_n
    outcomes = {}
    for r in rows:
        labs = outcomes.setdefault(r.get("outcome", "").lower(), [])
        if r.get("label", "") not in labs:
            labs.append(r.get("label", ""))
    for r in rows:
        imp, sat = float(r["importance"]), float(r["satisfaction"])
        r["importance"], r["satisfaction"] = imp, sat
        r["segment"] = (r.get("segment") or r.get("persona") or "All") if by_persona else "All personas"
        r["opportunity"] = round(opportunity(imp, sat), 2)
        r["zone"] = zone(imp, sat)
        flags = []
        if r.get("n") and jf.norm_source(r.get("source") or "survey") == "survey" and not assumed_flag:
            n = int(float(r["n"]))
            ni, ns = r.get("_n") or (n, n)
            r["margin"] = f"±{max(margin(imp, ni), margin(sat, ns)):.0f}"
            r["zone_stability"] = stability(imp, sat, ni, r["zone"], n_sat=ns)
            if r.get("_scored_on"):
                flags.append(f"scored on {r['_scored_on']} respondents only")
            if r.get("_b2b"):
                bi, bs = r["_b2b"]
                if imp >= 30 and bi >= 30:
                    flags.append(f"split importance: {imp:.0f}% rate it 4-5, {bi:.0f}% rate it 1-2")
                if bs >= 50:
                    flags.append(f"strong pain: {bs:.0f}% rate satisfaction 1-2")
            if n < MIN_N:
                flags.append(f"small sample (n={n} < {MIN_N})")
            sp = r.get("_split")
            if sp and len(sp) > 1:
                small = [f"{p} n={v[0]}" for p, v in sp.items() if v[0] < MIN_N]
                if small:
                    flags.append("small persona sample: " + ", ".join(small))
                imps = sorted(sp.items(), key=lambda kv: kv[1][1])
                if imps[-1][1][1] - imps[0][1][1] >= 20:
                    flags.append(f"personas disagree on importance: {imps[-1][0]} {imps[-1][1][1]} vs {imps[0][0]} {imps[0][1][1]}")
                r["persona_split"] = "; ".join(f"{p}: n={v[0]}, imp {v[1]}, sat {v[2]}" for p, v in sp.items())
            if r["zone_stability"] < STABLE:
                flags.append(f"zone could change within the margin of error ({r['zone_stability']}% of resamples agree)")
        else:
            r["margin"] = ""
            opp = r["opportunity"]
            if (any(abs(opp - t) <= 0.25 for t in (SOLID_OPP, HIGH_OPP))
                    or any(abs(sat - t) <= 1 for t in (MODEST_SAT_MIN, TABLE_STAKES_SAT))):
                flags.append("borderline")
        barrier, goal = r.get("barrier", ""), r.get("goal", "")
        if sat >= MODEST_SAT_MIN and NOT_EXIST.search(barrier):
            flags.append("barrier says it can't be done, yet satisfaction ≥ 80")
        if SOLUTION.search(goal):
            flags.append("possible solution in goal")
        if COMPANY.search(barrier + " " + goal + " " + r.get("context", "")):
            flags.append("company voice / jargon")
        if re.search(r"\band\b.*\b(chat|communicat|also)\b|\n", barrier + " " + goal, re.I):
            flags.append("may bundle two needs")
        if len(outcomes.get(r.get("outcome", "").lower(), [])) > 2:
            flags.append("outcome shared by %d statements" % len(outcomes[r.get('outcome', '').lower()]))
        if r.get("_no_survey"):
            flags.insert(0, "no survey answers: AI estimate kept")
        r["flags"] = "; ".join(flags)
    personas = {}
    for r in rows:
        personas.setdefault(r["segment"], set()).add(r.get("persona") or "?")
    for r in rows:
        ps = personas[r["segment"]]
        if by_persona and r["segment"] != "All (combined)" and (len(ps) > 1 or "?" in ps):
            r["flags"] = "; ".join(filter(None, [r["flags"], "persona missing or mixed in segment"]))
    rows.sort(key=lambda r: (r["segment"], -r["opportunity"]))
    return rows, assumed


COLS = ["segment", "label", "persona", "stage", "context", "barrier", "goal", "outcome",
        "n", "importance", "satisfaction", "source", "margin", "opportunity", "zone", "zone_stability",
        "flags", "persona_split", "evidence", "original_context", "original_barrier", "original_goal", "original_outcome",
        "edit_reason", "survey_importance", "survey_satisfaction"]


def you(text):
    """First person -> second person for survey wording ("When I review my" -> "When you review your")."""
    for a, b in [(r"\bI am\b", "you are"), (r"\bI was\b", "you were"), (r"\bI'm\b", "you're"), (r"\bI've\b", "you've"), (r"\bI\b", "you"), (r"\bmy\b", "your"),
                 (r"\bMy\b", "Your"), (r"\bme\b", "you"), (r"\bmyself\b", "yourself")]:
        text = re.sub(a, b, text)
    return text


def job_id(label):
    return f"J{int(label):02d}" if re.fullmatch(r"\d+", str(label or "")) else str(label or "")


def survey_item(r):
    """(job_id, job statement, context, importance question, satisfaction question), as the survey kit builds them."""
    ctx = you(r.get("context", "").strip().rstrip(",. "))
    goal = re.sub(r"^help me\s+", "", r.get("goal", "").strip().rstrip("."), flags=re.I)
    goal = you(goal)
    return (job_id(r.get("label")), goal, ctx,
            f"{ctx}, how important is it to you to {goal}?",
            f"{ctx}, how satisfied are you with how well you can {goal} today?")


def write_csv(rows, path):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, COLS, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


XL_HEAD = ["Label", "Persona", "Stage", "Context", "Barrier", "Goal", "Outcome", "Importance",
           "Satisfaction", "Opportunity score", "Zone", "Score source", "Evidence confidence", "n",
           "Zone stability %", "Review flags", "Importance reason", "Satisfaction reason"]
SRC_LABEL = {"ai": "AI estimate", "evidence_update": "Evidence update", "judgement": "Your judgement",
             "survey": "Survey"}


def write_xlsx(rows, path, assumed, logs=()):
    """Opportunity score and zone are live formulas that read the Rules sheet."""
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Opportunity scores"
    ws.append(XL_HEAD)
    for c in ws[1]:
        c.font = Font(bold=True)
    for i, r in enumerate(rows, start=2):
        src = jf.norm_source(r.get("source") or ("assumed" if assumed else "survey"))
        ws.append([r.get("label"), r.get("persona"), r.get("stage"), r.get("context"), r.get("barrier"),
                   r.get("goal"), r.get("outcome"), r["importance"], r["satisfaction"],
                   f"=H{i}/10+MAX(H{i}-I{i},0)/10",
                   f'=IF(I{i}>H{i},"Over-served",IF(I{i}>Rules!$B$2,"Table stakes",'
                   f'IF(AND(I{i}>=Rules!$B$3,H{i}>=Rules!$B$4),"Modest opportunity",'
                   f'IF(AND(I{i}<Rules!$B$3,J{i}>Rules!$B$6),"High opportunity",'
                   f'IF(AND(I{i}<Rules!$B$3,J{i}>=Rules!$B$5),"Solid opportunity","Served right")))))',
                   SRC_LABEL.get(src, src) + (" (over measured data)" if src == "judgement" and r.get("survey_importance") else ""),
                   r.get("confidence", ""), r.get("n", ""),
                   r.get("zone_stability", ""), r.get("flags", ""),
                   r.get("importance_reason", ""), r.get("satisfaction_reason", "")])
        ws[f"J{i}"].number_format = "0.0"
    for col, w in zip("ABCDEFGHIJKLMNOPQR", [7, 16, 11, 32, 36, 38, 32, 11, 12, 12, 18, 12, 12, 6, 10, 36, 40, 40]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    edit = PatternFill("solid", fgColor=tokens()["light"]["accent-tint"].lstrip("#"))
    for row in ws.iter_rows(min_row=2, min_col=8, max_col=9):
        for c in row:
            c.fill = edit
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:R{len(rows) + 1}"
    if assumed:
        ws.insert_rows(1)
        ws["A1"] = (jf.based_on(rows) + ". " + jf.confidence_note(rows) + " Edit the blue "
                    "Importance and Satisfaction cells: score and zone recalculate.")
        ws["A1"].font = Font(bold=True, color=tokens()["light"]["warn-ink"].lstrip("#"))
        # formulas shifted by insert_rows are not rewritten by openpyxl; rebuild them
        for i in range(3, len(rows) + 3):
            ws[f"J{i}"] = f"=H{i}/10+MAX(H{i}-I{i},0)/10"
            ws[f"K{i}"] = (f'=IF(I{i}>H{i},"Over-served",IF(I{i}>Rules!$B$2,"Table stakes",'
                           f'IF(AND(I{i}>=Rules!$B$3,H{i}>=Rules!$B$4),"Modest opportunity",'
                           f'IF(AND(I{i}<Rules!$B$3,J{i}>Rules!$B$6),"High opportunity",'
                           f'IF(AND(I{i}<Rules!$B$3,J{i}>=Rules!$B$5),"Solid opportunity","Served right")))))')
        ws.freeze_panes = "B3"
        ws.auto_filter.ref = f"A2:R{len(rows) + 2}"

    ev = wb.create_sheet("Evidence")
    ev.append(["Label", "Kind", "Source ID", "Excerpt"])
    for r in rows:
        for kind, key in (("Supports", "evidence"), ("Points the other way", "contradicts")):
            for e in split_evidence(r.get(key, "")):
                ev.append([r.get("label"), kind, e["id"], e["text"]])
    for col, w in zip("ABCD", [7, 20, 10, 90]):
        ev.column_dimensions[col].width = w

    sv = wb.create_sheet("Survey")
    sv.append(["job_id", "Job statement (help me …)", "Context (When you …)", "Persona", "Stage",
               "Importance question (1-5)", "Satisfaction question (1-5)", "Review before sending"])
    seen = set()
    for r in sorted(rows, key=lambda r: (len(str(r.get("label"))), str(r.get("label")))):
        if r.get("label") in seen:
            continue
        seen.add(r.get("label"))
        jid, goal, ctx, qi, qs = survey_item(r)
        sv.append([jid, goal, ctx, r.get("persona"), r.get("stage"), qi, qs, r.get("flags", "")])
    sv.append([])
    sv.append(["Columns A–E paste into the survey kit's Job key sheet. Ratings come back as J01_imp, J01_sat … "
               "and load with --ratings. Check the wording reads naturally before sending."])
    for c in list(ev[1]) + list(sv[1]):
        c.font = Font(bold=True)
    for col, w in zip("ABCDEFGH", [7, 40, 34, 16, 12, 60, 60, 30]):
        sv.column_dimensions[col].width = w

    lg = wb.create_sheet("Decision log")
    lg.append(["When (UTC)", "Job", "Changed", "From", "To", "Zone before", "Zone after", "Source", "By", "Reason", "Evidence"])
    for g in logs:
        lg.append([g.get("logged_at"), g.get("label"), g.get("field"), g.get("from"), g.get("to"), g.get("zone_from"),
                   g.get("zone_to"), "decision" if g.get("field") == "decision" else SRC_LABEL.get(jf.norm_source(g.get("source")), g.get("source")), g.get("by"),
                   g.get("reason"), g.get("evidence_ids")])
    if not logs:
        lg.append(["", "", "No changes yet. Every change to a score or wording is logged here, append-only."])
    for c in lg[1]:
        c.font = Font(bold=True)
    for col, w in zip("ABCDEFGHIJK", [20, 6, 12, 10, 10, 18, 18, 16, 8, 50, 16]):
        lg.column_dimensions[col].width = w

    rl = wb.create_sheet("Rules")
    for row in [["Rule", "Value", "Meaning"],
                ["Table stakes: satisfaction above", TABLE_STAKES_SAT, "Over-served is checked first (satisfaction > importance)"],
                ["Modest opportunity: satisfaction from", MODEST_SAT_MIN, "...up to the table-stakes value, with importance at least the next row"],
                ["Modest opportunity: importance at least", MODEST_IMP_MIN, ""],
                ["Solid opportunity: score from", SOLID_OPP, "With satisfaction below the Modest minimum"],
                ["High opportunity: score above", HIGH_OPP, "With satisfaction below the Modest minimum"],
                ["Opportunity score", "", "importance/10 + max(importance - satisfaction, 0)/10"]]:
        rl.append(row)
    rl["A1"].font = rl["B1"].font = rl["C1"].font = Font(bold=True)
    for col, w in zip("ABC", [34, 8, 70]):
        rl.column_dimensions[col].width = w
    wb.save(path)


def split_evidence(text):
    out = []
    for part in (text or "").split("||"):
        part = part.strip()
        if not part:
            continue
        m = re.match(r"^([A-Z]{2,5}-[\w-]+)\s*:\s*(.*)$", part, re.S)
        out.append({"id": m.group(1), "text": m.group(2).strip()} if m else {"id": "", "text": part})
    return out


def tokens():
    """Colour tokens shared by the page, its image export and the chart image (templates/tokens.json)."""
    import json
    p = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "templates", "tokens.json")
    return json.load(open(p, encoding="utf-8"))


def css_vars(t):
    return " ".join(f"--{k}:{v};" for k, v in t.items())


def write_html(rows, path, title, study, logs=(), meta=None):
    """Fill templates/landscape.html with the scored jobs."""
    import json, datetime
    tpl = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "templates", "landscape.html")
    page = open(tpl, encoding="utf-8").read()
    ids = {e["id"] for r in rows for k in ("evidence", "contradicts") for e in split_evidence(r.get(k, "")) if e["id"]}
    kinds = [("INT", "interview", "interviews"), ("FD", "staff interview", "staff interviews"),
             ("SUP", "support item", "support items"),
             ("FUN", "funnel observation", "funnel observations"), ("DOC", "document", "documents"),
             ("SUR", "survey note", "survey notes")]
    parts = []
    for pre, one, many in kinds:
        n = len({i.split("-")[0] + "-" + i.split("-")[1] for i in ids if i.startswith(pre + "-")})
        if n:
            parts.append(f"{n} {one if n == 1 else many}")
    data = {
        "title": title, "study": study, "generated": datetime.date.today().isoformat(), "version": VERSION,
        "meta": meta or {}, "based_on": jf.based_on(rows), "confidence_note": jf.confidence_note(rows),
        "logs": [{k: g.get(k, "") for k in ("label", "logged_at", "field", "from", "to", "zone_from", "zone_to",
                                            "source", "by", "reason", "evidence_ids", "version")} for g in logs],
        "sources_summary": ", ".join(parts) or "the research provided",
        "rules": {"tableStakesSat": TABLE_STAKES_SAT, "underSatMin": MODEST_SAT_MIN, "underImpMin": MODEST_IMP_MIN,
                  "solidOpp": SOLID_OPP, "highOpp": HIGH_OPP},
        "jobs": [{
            "label": r.get("label"), "persona": r.get("persona") or "Unassigned", "stage": r.get("stage", ""),
            "context": r.get("context", ""), "barrier": r.get("barrier", ""), "goal": r.get("goal", ""),
            "outcome": r.get("outcome", ""), "importance": r["importance"], "satisfaction": r["satisfaction"],
            "n": r.get("n", ""), "source": jf.norm_source(r.get("source") or ("survey" if r.get("n") else "ai")),
            "evidence_ids": r.get("evidence_ids", ""), "provenance": r.get("provenance", ""),
            "confidence": r.get("confidence", ""), "importance_reason": r.get("importance_reason", ""),
            "satisfaction_reason": r.get("satisfaction_reason", ""), "flags": r.get("flags", ""),
            "persona_split": r.get("persona_split", ""),
            "edit_reason": r.get("edit_reason", ""),
            "survey_importance": r.get("survey_importance", ""), "survey_satisfaction": r.get("survey_satisfaction", ""),
            "original": {f: r.get("original_" + f, "") for f in ("context", "barrier", "goal", "outcome")}
                        if any(r.get("original_" + f) for f in ("context", "barrier", "goal", "outcome")) else None,
            "evidence": split_evidence(r.get("evidence", "")), "contradicts": split_evidence(r.get("contradicts", "")),
        } for r in rows]}
    blob = json.dumps(data, ensure_ascii=False).replace("</", "<\\/")
    tk = tokens()
    page = (page.replace("/*__LANDSCAPE_DATA__*/null", blob).replace("__PAGE_TITLE__", title)
            .replace("/*__TOKENS_LIGHT__*/", css_vars(tk["light"])).replace("/*__TOKENS_DARK__*/", css_vars(tk["dark"]))
            .replace("/*__TOKENS_JSON__*/null", json.dumps({"light": tk["light"]})))
    open(path, "w", encoding="utf-8").write(page)


def chart(rows, seg, path, assumed, title_prefix):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import Polygon, Rectangle
    T = tokens()["light"]
    fig, ax = plt.subplots(figsize=(12, 8), dpi=130, facecolor=T["surface"])
    ax.set_facecolor(T["plot"])
    ax.set_xlim(0, 100); ax.set_ylim(0, 100)
    # regions (team guide)
    ax.add_patch(Polygon([(0, 0), (100, 100), (0, 100)], color=T["z-over"], lw=0))
    ax.add_patch(Polygon([(TABLE_STAKES_SAT, TABLE_STAKES_SAT), (100, 100), (100, TABLE_STAKES_SAT)],
                         color=T["z-stakes"], lw=0))
    ax.add_patch(Rectangle((MODEST_IMP_MIN, MODEST_SAT_MIN), 100 - MODEST_IMP_MIN,
                           TABLE_STAKES_SAT - MODEST_SAT_MIN, color=T["z-under"], lw=0))
    # opportunity = SOLID_OPP  <=>  sat = 2*imp - 10*SOLID_OPP
    x0 = 5 * SOLID_OPP; x1 = (MODEST_SAT_MIN + 10 * SOLID_OPP) / 2
    ax.add_patch(Polygon([(x0, 0), (100, 0), (100, MODEST_SAT_MIN), (x1, MODEST_SAT_MIN)],
                         color=T["z-solid"], lw=0))
    h0 = 5 * HIGH_OPP  # opportunity = HIGH_OPP <=> sat = 2*imp - 10*HIGH_OPP
    ax.add_patch(Polygon([(h0, 0), (100, 0), (100, min(MODEST_SAT_MIN, 200 - 10 * HIGH_OPP))],
                         color=T["z-high"], lw=0))
    ax.grid(color=T["ink"], alpha=.07)
    import math
    bold = dict(fontweight="bold", zorder=1.5)
    ax.text(4, 93, "Over-served", fontsize=12, color=T["z-over-ink"], **bold)
    ax.text(4, 3, "Served right", fontsize=12, color=T["muted"], **bold)
    ax.text(98, 3, "Underserved", fontsize=13, color=T["z-high-ink"], ha="right", **bold)
    # band labels along each band (one point of satisfaction = half a point of importance)
    ang = math.degrees(math.atan2(2, 1))  # in data space; transform_rotates_text keeps it after layout
    for score, label, ink in (((SOLID_OPP + HIGH_OPP) / 2 - .25, "Solid opportunity", T["z-solid-ink"]),
                              (HIGH_OPP + 1.6, "High opportunity", T["z-high-ink"])):
        ax.text((4 + 10 * score) / 2, 4, label, rotation=ang, rotation_mode="anchor", transform_rotates_text=True, fontsize=10, color=ink, **bold)
    # the two small zones are labelled outside the plot with a leader line, so markers can't cover them
    for (zx, zy), (lx, ly), label, ink in (((97, (TABLE_STAKES_SAT + 100) / 2 + 3), (104, 97), "Table stakes", T["ink"]),
                                           ((97, (TABLE_STAKES_SAT + MODEST_SAT_MIN) / 2), (104, MODEST_SAT_MIN - 1), "Modest\nopportunity", T["z-under-ink"])):
        ax.annotate(label, xy=(zx, zy), xytext=(lx, ly), annotation_clip=False, va="center", fontsize=10, color=ink,
                    fontweight="bold", arrowprops=dict(arrowstyle="-", color=T["muted"], lw=.8))
    pts = [r for r in rows if r["segment"] == seg]
    placed = []
    for r in pts:
        x, y = r["importance"], r["satisfaction"]
        src = jf.norm_source(r.get("source"))  # same marker rules as the page (references/score-sources.md)
        if src in ("ai", "evidence_update"):
            ring = T["z-high-ink"] if zone(x, y) == "High opportunity" else T["est-ink"]  # one marker, readable on dark green
            ax.scatter(x, y, s=220, zorder=2, facecolor="none", edgecolor=ring, lw=1.1, linestyle=(0, (2, 1.5)))
        ax.scatter(x, y, s=70, zorder=3, marker="s" if src == "judgement" else "o", facecolor=T["dot"],
                   edgecolor=T["dot-ink"], lw=1.4)
        dy = 2.2
        for (px, py) in placed:  # nudge labels that would collide
            if abs(px - x) < 2.5 and abs(py - (y + dy)) < 2.2:
                dy = -3.6 if dy > 0 else dy - 2.2
        placed.append((x, y + dy))
        ax.text(x, y + dy, str(r.get("label", "")), ha="center", fontsize=9, zorder=4,
                color=T["z-high-ink"] if zone(x, y + dy) == "High opportunity" else T["ink"])
    sub = jf.based_on(pts) + ("" if not assumed else " — estimates and judgement are not survey data")
    name = title_prefix.strip() or "JTBD opportunity map"
    ax.set_title(f"{name} · {seg}\n{sub}", fontsize=12)
    ax.set_xlabel("Importance (% rating 4–5)"); ax.set_ylabel("Satisfaction (% rating 4–5)")
    ax.text(0, -11, "Dashed ring = AI estimate from customer evidence · square = team judgement, not customer evidence · "
            "plain = measured by survey.\nHigh, Solid and Modest together are Underserved.",
            fontsize=8, color=T["muted"], transform=ax.transData)
    fig.tight_layout(); fig.savefig(path); plt.close(fig)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("input")
    ap.add_argument("--out", default=".")
    ap.add_argument("--name")
    ap.add_argument("--title", default="", help="prefix for chart titles, e.g. product name")
    ap.add_argument("--assumed", action="store_true", help="ratings are assumptions")
    ap.add_argument("--ratings", help="respondent-level ratings CSV (1-5 scale)")
    ap.add_argument("--pool", action="store_true",
                    help="with --ratings, score every job on all respondents (default: the job's own persona)")
    ap.add_argument("--combined", action="store_true", help="with --by-persona and --ratings, add an all-persona view")
    ap.add_argument("--by-persona", action="store_true", help="one chart per persona (default: one view of all personas)")
    ap.add_argument("--html", action="store_true", help="also write the interactive page")
    ap.add_argument("--study", default="", help="study name shown on the page")
    ap.add_argument("--question", default="", help="the research question (about people)")
    ap.add_argument("--decision", default="", help="the decision it informs (read-out only)")
    ap.add_argument("--framing-note", default="", help="skew note when the user kept a product-first question")
    ap.add_argument("--entry", default="", help="where the run started: research, quotes, jobs or survey")
    a = ap.parse_args()
    name = a.name or os.path.splitext(os.path.basename(a.input))[0]
    os.makedirs(a.out, exist_ok=True)
    rows = read(a.input) if not a.ratings else None
    if a.ratings:
        rows = aggregate(read_statements(a.input), a.ratings, a.combined, a.by_persona, a.pool)
        rated = {r.get("label") for r in rows}
        try:
            estimates = read(a.input)
        except (StopIteration, KeyError):
            estimates = []
        kept = [r for r in estimates if r.get("label") not in rated]
        for r in kept:  # no survey answers: keep the AI estimate, labelled as one
            r["source"] = r.get("source") or "ai"
            r["_no_survey"] = True
        if kept and not a.by_persona:
            print(f"{len(kept)} job(s) had no survey answers and keep their AI estimate: "
                  + ", ".join(str(r.get("label")) for r in kept))
            rows += kept
    rows, assumed = score(rows, a.assumed, a.by_persona)
    logs = read_logs(a.input)
    built = []  # (ok, what, file or reason)
    write_csv(rows, os.path.join(a.out, f"{name}_ranked.csv"))
    built.append((True, "Ranked CSV", f"{name}_ranked.csv"))
    try:
        write_xlsx(rows, os.path.join(a.out, f"{name}_scores.xlsx"), assumed, logs)
        built.append((True, "Spreadsheet", f"{name}_scores.xlsx"))
    except ImportError:
        built.append((False, "Spreadsheet", "openpyxl isn't installed. Fix: pip install openpyxl"))
    try:
        for seg in sorted({r["segment"] for r in rows}):
            safe = re.sub(r"[^A-Za-z0-9]+", "_", seg).strip("_")
            fn = f"{name}_map.png" if not a.by_persona else f"{name}_{safe}.png"
            chart(rows, seg, os.path.join(a.out, fn), assumed, a.title + " " if a.title else "")
            built.append((True, "Chart image", fn))
    except ImportError:
        built.append((False, "Chart image", "matplotlib isn't installed. Fix: pip install matplotlib"))
    if a.html:
        write_html(rows, os.path.join(a.out, f"{name}_map.html"), a.title or name, a.study, logs,
                   {"question": a.question, "decision": a.decision, "framing_note": a.framing_note, "entry": a.entry})
        built.append((True, "Interactive page", f"{name}_map.html"))
    print(f"{jf.based_on(rows)}; {len(rows)} jobs. {jf.confidence_note(rows)}")
    if a.framing_note:
        print("Framing note: " + a.framing_note)
    for seg in dict.fromkeys(r["segment"] for r in rows):  # use these counts in the read-out
        z = [r["zone"] for r in rows if r["segment"] == seg]
        n = {k: z.count(k) for k in ("High opportunity", "Solid opportunity", "Modest opportunity",
                                     "Table stakes", "Served right", "Over-served")}
        under = n["High opportunity"] + n["Solid opportunity"] + n["Modest opportunity"]
        print(f"Zones{'' if seg in ('', 'All', 'All personas') else ' [' + seg + ']'}: Underserved {under} of {len(z)} "
              f"(High {n['High opportunity']}, Solid {n['Solid opportunity']}, Modest {n['Modest opportunity']}) · "
              f"Table stakes {n['Table stakes']} · Served right {n['Served right']} · Over-served {n['Over-served']}")
    for r in rows:
        print(f"  [{r['segment']}] {r.get('label', ''):>3}  imp {r['importance']:5.1f}  "
              f"sat {r['satisfaction']:5.1f}  opp {r['opportunity']:5.2f}  {r['zone']:<18} {r['flags']}")
    print("\nFiles in " + os.path.abspath(a.out) + ":")
    for ok, what, info in built:
        print(f"  {'[x]' if ok else '[ ] NOT BUILT'} {what}: {info}")
    if not all(ok for ok, _, _ in built):
        print("Tell the user which files are missing and how to fix it; don't present the run as complete.")


if __name__ == "__main__":
    main()
